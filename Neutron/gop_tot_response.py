#!/usr/bin/env python3
"""
Gộp giá trị "Tot. response (p/cmq/pr)" từ nhiều file *_FLUKA_<n>_sum.lis
(n chạy từ 21 đến 65) thành 1 file Excel duy nhất.

Cách dùng:
    python gop_tot_response.py                # quét thư mục hiện tại
    python gop_tot_response.py /duong/dan/thu/muc
    python gop_tot_response.py /duong/dan/thu/muc ket_qua.xlsx

Script tự tìm file có tên chứa "FLUKA_<n>_sum.lis" (không cần biết
phần tiền tố phía trước, ví dụ "1784207410693_phits_ca_fission_FLUKA_21_sum.lis"
hay "phits_ca_fission_FLUKA_21_sum.lis" đều được nhận).
Lấy cả GIÁ TRỊ và SAI SỐ (+/- %) của "Tot. response".
"""

import glob
import os
import re
import sys

import openpyxl
from openpyxl.styles import Font

START_IDX = 21
END_IDX = 65

# Regex bắt giá trị số VÀ sai số (%) ngay sau "Tot. response (p/cmq/pr)"
# Ví dụ dòng gốc:
#   Tot. response (p/cmq/pr)  2.7086708E-05  +/-  0.6821581     %
TOT_RESPONSE_RE = re.compile(
    r"Tot\.\s*response\s*\(p/cmq/pr\)\s*"
    r"([-+0-9.EeDd]+)\s*"          # giá trị
    r"\+/-\s*"
    r"([-+0-9.EeDd]+)\s*%"          # sai số (%)
)


def find_file_for_index(folder, idx):
    """Tìm file có tên chứa FLUKA_<idx>_sum.lis trong thư mục (không phân biệt tiền tố)."""
    pattern = os.path.join(folder, f"*FLUKA_{idx}_sum.lis")
    matches = glob.glob(pattern)
    if not matches:
        return None
    # nếu có nhiều hơn 1 kết quả, lấy file đầu tiên (sắp xếp cho ổn định)
    matches.sort()
    return matches[0]


def extract_tot_response(filepath):
    """Đọc file và trả về (giá trị, sai số %) dạng float, hoặc (None, None) nếu không tìm thấy."""
    with open(filepath, "r", errors="ignore") as f:
        content = f.read()
    m = TOT_RESPONSE_RE.search(content)
    if not m:
        return None, None

    raw_value = m.group(1).replace("D", "E").replace("d", "e")
    raw_err = m.group(2).replace("D", "E").replace("d", "e")

    try:
        value = float(raw_value)
    except ValueError:
        value = None

    try:
        err = float(raw_err)
    except ValueError:
        err = None

    return value, err


def main():
    folder = sys.argv[1] if len(sys.argv) > 1 else "."
    out_path = sys.argv[2] if len(sys.argv) > 2 else "tong_hop_Tot_response.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tot response"

    header = ["Index", "Tên file", "Tot. response (p/cmq/pr)", "Sai so (+/- %)"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, name="Arial")

    missing = []
    for idx in range(START_IDX, END_IDX + 1):
        fpath = find_file_for_index(folder, idx)
        if fpath is None:
            missing.append(idx)
            ws.append([idx, "KHONG TIM THAY FILE", None, None])
            continue

        value, err = extract_tot_response(fpath)
        if value is None:
            missing.append(idx)
            ws.append([idx, os.path.basename(fpath), "KHONG DOC DUOC GIA TRI", None])
            continue

        ws.append([idx, os.path.basename(fpath), value, err])

    # font Arial cho toàn bộ dữ liệu
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial")

    # định dạng khoa học cho cột giá trị
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000E+00"

    # định dạng số thập phân cho cột sai số (%)
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"

    # auto width
    for col_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = length + 4

    wb.save(out_path)

    print(f"Da luu ket qua vao: {out_path}")
    if missing:
        print(f"Canh bao: khong lay duoc du lieu cho cac index: {missing}")
    else:
        print("Da lay du du lieu cho tat ca index tu 21 den 65.")


if __name__ == "__main__":
    main()
